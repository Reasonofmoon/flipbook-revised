"""Branded sample deck for the READMASTER flipbook landing-page GIFs.

Original content (written for this demo): a short English reading lesson
with Korean glosses, laid out like an academy workbook.
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

INK = RGBColor(0x1C, 0x2B, 0x4B)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
PENCIL = RGBColor(0x6A, 0x71, 0x80)
RULE = RGBColor(0xE3, 0xE6, 0xEC)
HIGHLIGHT = RGBColor(0xFF, 0xE3, 0x4D)
REDPEN = RGBColor(0xD6, 0x45, 0x3D)
FONT = "Noto Sans CJK KR"

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, lines, size=24, color=INK, bold=False, align=PP_ALIGN.LEFT, spacing=1.15):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tb


def page_frame(slide, unit, page):
    box(slide, 0, 0, W, H, fill=PAPER)
    text(slide, Emu(560000), Emu(330000), Emu(6000000), Emu(400000),
         f"READMASTER Reading Lab  |  {unit}", size=13, color=PENCIL)
    text(slide, W - Emu(1500000), Emu(330000), Emu(940000), Emu(400000),
         str(page), size=13, color=PENCIL, align=PP_ALIGN.RIGHT)
    box(slide, Emu(560000), Emu(760000), W - Emu(1120000), Emu(12000), fill=RULE)


def highlight_under(slide, x, y, w):
    box(slide, x, y, w, Emu(150000), fill=HIGHLIGHT)


# 1. Cover
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, W, H, fill=INK)
box(s, Emu(800000), Emu(2950000), Emu(5200000), Emu(230000), fill=HIGHLIGHT)
text(s, Emu(800000), Emu(1250000), Emu(9000000), Emu(500000), "READMASTER Reading Lab", size=20, color=HIGHLIGHT, bold=True)
text(s, Emu(800000), Emu(1850000), Emu(10000000), Emu(1300000), "Unit 3. The Honeybee Dance", size=54, color=PAPER, bold=True)
text(s, Emu(800000), Emu(3500000), Emu(10000000), Emu(900000), ["꿀벌은 어떻게 꽃밭의 위치를 알려 줄까?", "중2–중3 · 독해 + 핵심 어휘 6개"], size=22, color=PAPER)
text(s, Emu(800000), Emu(5700000), Emu(8000000), Emu(500000), "리드마스터 영어 · 2026 가을 학기", size=14, color=RGBColor(0xB8, 0xC2, 0xD6))

# 2. Warm-up
s = prs.slides.add_slide(BLANK)
page_frame(s, "Unit 3", 2)
text(s, Emu(560000), Emu(1050000), Emu(9000000), Emu(700000), "Warm-up", size=36, bold=True)
qs = ["1. 꿀벌 한 마리가 하루에 방문하는 꽃은 몇 송이일까요?",
      "2. 말을 하지 못하는 동물은 어떻게 정보를 나눌까요?",
      "3. 'dance'라는 단어가 과학 글에 나오면 어떤 뜻일까요?"]
for i, q in enumerate(qs):
    y = Emu(2000000 + i * 1150000)
    box(s, Emu(560000), y, W - Emu(1120000), Emu(950000), fill=None, line=RULE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, Emu(860000), y + Emu(230000), Emu(10000000), Emu(600000), q, size=24)

# 3. Passage
s = prs.slides.add_slide(BLANK)
page_frame(s, "Unit 3", 3)
text(s, Emu(560000), Emu(1000000), Emu(9000000), Emu(700000), "Reading", size=36, bold=True)
highlight_under(s, Emu(560000), Emu(2520000), Emu(5350000))
passage = [
    "When a honeybee finds a field of flowers, it flies back to the hive",
    "and performs a waggle dance. The angle of the dance shows the",
    "direction of the flowers, and the length of the waggle shows how far",
    "away they are. Other bees watch closely, then fly straight to the food.",
]
text(s, Emu(560000), Emu(1950000), Emu(11000000), Emu(3000000), passage, size=24, spacing=1.6)
text(s, Emu(560000), Emu(5450000), Emu(11000000), Emu(500000), "✎ 형광펜: 춤의 각도 = 방향, 흔드는 길이 = 거리", size=18, color=REDPEN)

# 4. Key vocabulary (table)
s = prs.slides.add_slide(BLANK)
page_frame(s, "Unit 3", 4)
text(s, Emu(560000), Emu(1000000), Emu(9000000), Emu(700000), "핵심 어휘", size=36, bold=True)
words = [("hive", "벌집"), ("perform", "수행하다, 보여 주다"), ("angle", "각도"),
         ("direction", "방향"), ("length", "길이"), ("closely", "자세히, 면밀히")]
rows, cols = len(words) + 1, 2
tbl = s.shapes.add_table(rows, cols, Emu(560000), Emu(1900000), Emu(7600000), Emu(4200000)).table
tbl.columns[0].width, tbl.columns[1].width = Emu(3200000), Emu(4400000)
for c, head in enumerate(["English", "뜻"]):
    cell = tbl.cell(0, c)
    cell.text = head
for r, (en, ko) in enumerate(words, start=1):
    tbl.cell(r, 0).text, tbl.cell(r, 1).text = en, ko
for r in range(rows):
    for c in range(cols):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = INK if r == 0 else (PAPER if r % 2 else RGBColor(0xF4, 0xF6, 0xF9))
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.name = FONT
                run.font.size = Pt(20)
                run.font.bold = r == 0
                run.font.color.rgb = PAPER if r == 0 else INK
box(s, Emu(8600000), Emu(1900000), Emu(3000000), Emu(4200000), fill=HIGHLIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Emu(8850000), Emu(2150000), Emu(2500000), Emu(3800000),
     ["Tip", "", "angle과 length는", "이 글의 정답 단서예요.", "", "본문에서 두 단어를", "찾아 동그라미!"], size=19, bold=False)

# 5. Check-up quiz
s = prs.slides.add_slide(BLANK)
page_frame(s, "Unit 3", 5)
text(s, Emu(560000), Emu(1000000), Emu(9000000), Emu(700000), "Check-up", size=36, bold=True)
text(s, Emu(560000), Emu(1900000), Emu(11000000), Emu(700000),
     "Q. 꿀벌 춤에서 '흔드는 길이(the length of the waggle)'가 알려 주는 것은?", size=24, bold=True)
opts = ["① 꽃의 색깔", "② 꽃까지의 거리", "③ 꽃의 개수", "④ 벌집의 크기"]
for i, o in enumerate(opts):
    y = Emu(2850000 + i * 720000)
    if i == 1:
        box(s, Emu(560000), y, Emu(4600000), Emu(600000), fill=HIGHLIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, Emu(800000), y + Emu(80000), Emu(6000000), Emu(500000), o, size=24)
text(s, Emu(5700000), Emu(3600000), Emu(5600000), Emu(600000), "정답 ② — 본문 3번째 줄", size=20, color=REDPEN)

# 6. Summary
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, W, H, fill=INK)
text(s, Emu(800000), Emu(1100000), Emu(10000000), Emu(800000), "오늘의 한 줄 정리", size=24, color=HIGHLIGHT, bold=True)
text(s, Emu(800000), Emu(2000000), Emu(10600000), Emu(2000000),
     ["Bees dance to share directions:", "angle tells where, length tells how far."], size=40, color=PAPER, bold=True, spacing=1.3)
text(s, Emu(800000), Emu(4600000), Emu(10000000), Emu(600000), "다음 시간: Unit 4. Why Leaves Change Color", size=20, color=RGBColor(0xB8, 0xC2, 0xD6))

prs.save(sys.argv[1])
print("saved", sys.argv[1], len(prs.slides), "slides")
